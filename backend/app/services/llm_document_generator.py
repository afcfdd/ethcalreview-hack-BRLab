"""
LLMベース書類生成サービス

Chain of Thoughtアプローチでアカデミックな研究倫理書類を生成します。
"""

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path
from typing import Dict, Any, List
from datetime import datetime

from app.services.llm_client import LLMClient
from app.logger import get_logger

logger = get_logger(__name__)


# 用語制約（全 generator 共通）: 研究に協力する人の呼称は必ず「研究対象者」または「参加者」に統一する。
# 注意: 旧来の呼称（健康状態を含意する語や、実験の語を含む被る側の呼称など）は一切使用しない。
# プロンプトに禁止語そのものを書くと生成文に伝播し得るため、ここでは肯定形で統一を指示する。
TERMINOLOGY_RULE = (
    "用語の制約：研究に協力する人を指す場合は、必ず「研究対象者」または「参加者」と表記すること。"
    "それ以外の呼称（健康状態を含意する旧来の語や、実験の語を含む対象側の旧来の呼称など）は一切使用しないこと。"
    "実験計画の用語も同様に、対象側の旧来語を含む計画名は使わず、"
    "『参加者内計画（参加者内要因）』『参加者間計画（参加者間要因）』のように『参加者』を用いて表記すること。"
    "実験を実施・進行する担当者は「実験実施者」と表記し、「研究者」とは書かないこと"
    "（ただし役職を指す「研究責任者」はそのまま用いてよい）。"
)

# アカデミックライティング基本ルール（実施計画書向け）
IMPLEMENTATION_PLAN_RULES = """
【実施計画書の執筆原則】

★★★ 最重要ルール ★★★
1. すべては「決定済み」の前提で書く
   - 「〜と考えられる」「〜の可能性がある」ではなく「〜である」「〜を行う」と断定する
   - 計画書は実施が決まった内容を記述するものである

2. 専門用語・テクニカルタームは使用しない
   - 使用する場合は必ず括弧内に平易な説明を付ける
   - 例: 「fMRI（脳の活動を画像で見る装置）」

3. 文体は「である調」で統一

4. 段落間の無駄な空白行は入れない
   - 各セクションは連続した文章として記述する

5. 避けるべき表現：
   - コロン（:）、波線（~~）、感嘆符（！）
   - 箇条書き記号の乱用
   - 「素晴らしい」「重要です」などの強調表現

6. 使用すべき表現：
   - 「本研究では〜を行う」「〜を明らかにする」
   - 「参加者は〜する」「実験者は〜を測定する」
   - 「〜に基づき」「〜を踏まえて」

7. 用語の統一：
   - 研究に協力する人は必ず「研究対象者」または「参加者」と表記する
   - それ以外の呼称（健康状態を含意する旧来の語や、実験の語を含む対象側の旧来の呼称など）は一切使用しない
"""


def build_implementation_plan_outline(is_questionnaire: bool = False) -> List[Dict[str, Any]]:
    """実施計画書の章立て（見出し構成）を公式参考様式 260127 に沿って構築する。

    公式参考様式（【記載例】01-2.実施計画書（参考様式）260127）の構成:
        1.課題名
        2.研究の概要（半ページ以内で簡略に記載）
        3.実験方法
            3-1.実験の目的
            3-2.実験参加者
            3-3.実験装置・実験タスク
            3-4.実験手順
        【アンケート調査の場合】
        3.アンケートの実施方法
            3-1.アンケートの目的
            3-2.研究対象者
            3-3.実施内容

    Returns:
        各セクションの dict（number, title, level, key）のリスト。
        level=1 が大見出し、level=2 が小見出し。
    """
    outline: List[Dict[str, Any]] = [
        {"number": "1", "title": "課題名", "level": 1, "key": "research_title"},
        {"number": "2", "title": "研究の概要", "level": 1, "key": "overview"},
    ]
    if is_questionnaire:
        outline += [
            {"number": "3", "title": "アンケートの実施方法", "level": 1, "key": None},
            {"number": "3-1", "title": "アンケートの目的", "level": 2, "key": "experiment_objective"},
            {"number": "3-2", "title": "研究対象者", "level": 2, "key": "participants"},
            {"number": "3-3", "title": "実施内容", "level": 2, "key": "procedures"},
        ]
    else:
        outline += [
            {"number": "3", "title": "実験方法", "level": 1, "key": None},
            {"number": "3-1", "title": "実験の目的", "level": 2, "key": "experiment_objective"},
            {"number": "3-2", "title": "実験参加者", "level": 2, "key": "participants"},
            {"number": "3-3", "title": "実験装置・実験タスク", "level": 2, "key": "equipment"},
            {"number": "3-4", "title": "実験手順", "level": 2, "key": "procedures"},
        ]
    return outline


def _is_questionnaire_study(context: Dict[str, Any]) -> bool:
    """研究方法・タイトルからアンケート調査主体かを簡易判定する。"""
    source = " ".join(
        str(context.get(key, ""))
        for key in ("research_title", "brief_description", "methodology")
    )
    devices = context.get("devices") or []
    has_device = any(str(d).strip() for d in devices)
    questionnaire_markers = ["アンケート", "質問紙", "調査票", "Webアンケート", "ウェブアンケート"]
    return (not has_device) and any(marker in source for marker in questionnaire_markers)


class LLMDocumentGenerator:
    """
    LLMベースの書類生成器

    Chain of Thought (CoT) アプローチで各セクションを順次生成し、
    アカデミックライティングの原則に準拠した文章を生成します。
    """
    
    def __init__(self, llm_client: LLMClient, lab_defaults: Dict[str, Any]):
        self.llm = llm_client
        self.defaults = lab_defaults
    
    async def generate_implementation_plan(
        self,
        form_data: Dict[str, Any],
        output_dir: Path
    ) -> Path:
        """
        実施計画書をLLMで生成（公式参考様式 260127 の構成に準拠）

        構成（実験等の場合）:
            1.課題名 / 2.研究の概要（半ページ） /
            3.実験方法（3-1.実験の目的 / 3-2.実験参加者 / 3-3.実験装置・実験タスク / 3-4.実験手順）
        アンケート調査の場合は 3.アンケートの実施方法（3-1.目的 / 3-2.研究対象者 / 3-3.実施内容）に分岐。

        Chain of Thoughtで各セクションを順次生成する。
        """
        logger.info("=" * 60)
        logger.info("LLM実施計画書生成 開始")

        # コンテキスト準備
        context = {
            "research_title": form_data.get("title", form_data.get("research_title", "")),
            "brief_description": form_data.get("purpose", form_data.get("research_purpose", "")),
            "methodology": form_data.get("methodology", form_data.get("research_method", "")),
            "target_participants": form_data.get("targetDescription", ""),
            "duration": form_data.get("duration", form_data.get("duration_minutes", 0)),
            "participant_count": form_data.get("expectedParticipants", form_data.get("participant_count", 30)),
            "devices": form_data.get("devices", []),
            "risks": form_data.get("risks", []),
            "risk_countermeasures": form_data.get("riskCountermeasures", []),
            "reward_amount": form_data.get("rewardAmount", form_data.get("reward_amount", 1230)),
        }

        is_questionnaire = _is_questionnaire_study(context)
        outline = build_implementation_plan_outline(is_questionnaire)
        logger.info(f"  研究タイトル: {context['research_title'][:50]}...")
        logger.info(f"  様式分岐: {'アンケート調査' if is_questionnaire else '実験等'}")

        # 中間ファイル保存用
        intermediate_file = output_dir / "_intermediate_plan.json"

        # Chain of Thought: 各セクションを順次生成
        sections = {}

        # 2. 研究の概要（半ページ以内）
        logger.info("  [1] 研究の概要 生成中...")
        sections["overview"] = await self._generate_overview(context)
        self._save_intermediate(intermediate_file, sections, "overview完了")

        # 3-1. 実験の目的 / アンケートの目的
        logger.info("  [2] 目的 生成中...")
        sections["experiment_objective"] = await self._generate_experiment_objective(context, is_questionnaire)
        self._save_intermediate(intermediate_file, sections, "experiment_objective完了")

        # 3-2. 実験参加者 / 研究対象者
        logger.info("  [3] 研究対象者 生成中...")
        sections["participants"] = await self._generate_participants(context, is_questionnaire)
        self._save_intermediate(intermediate_file, sections, "participants完了")

        if not is_questionnaire:
            # 3-3. 実験装置・実験タスク
            logger.info("  [4] 実験装置・実験タスク 生成中...")
            sections["equipment"] = await self._generate_equipment(context)
            self._save_intermediate(intermediate_file, sections, "equipment完了")

        # 3-3/3-4. 実験手順 / 実施内容
        logger.info("  [5] 手順・実施内容 生成中...")
        sections["procedures"] = await self._generate_procedures(context, is_questionnaire)
        self._save_intermediate(intermediate_file, sections, "procedures完了")

        # DOCXファイル生成
        output_path = self._build_implementation_plan_docx(
            sections=sections,
            context=context,
            outline=outline,
            output_dir=output_dir
        )
        
        logger.info(f"LLM実施計画書生成 完了: {output_path.name}")
        logger.info("=" * 60)
        
        return output_path
    
    def _save_intermediate(self, path: Path, sections: Dict[str, str], status: str):
        """中間結果を保存"""
        import json
        data = {
            "status": status,
            "timestamp": datetime.now().isoformat(),
            "sections": sections
        }
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        logger.info(f"    中間保存: {status}")
    
    async def _generate_overview(self, context: Dict[str, Any]) -> str:
        """研究の概要を生成（公式参考様式「2.研究の概要」＝半ページ以内で簡略に）"""
        prompt = f"""
{IMPLEMENTATION_PLAN_RULES}

【タスク】
公式参考様式の「2.研究の概要」セクションを執筆してください。
これは半ページ以内で、研究の背景・目的・方法・研究対象者の概略を簡潔にまとめるものです。

研究タイトル: {context['research_title']}
研究概要: {context['brief_description']}
研究方法: {context['methodology']}
研究対象者: {context['target_participants']}

【記述すべき内容（簡略にまとめる）】
1. 研究の背景と、本研究が取り組む課題（2-3文）
2. 本研究の目的（何を明らかにするか）（1-2文）
3. どのような研究対象者に、どのような方法で実施するかの概略（2-3文）

【執筆スタイル】
- 半ページ以内（300-450文字程度）で簡略に記述する
- 「本研究では〜を明らかにする」のような断定形
- 段落間の空行は入れない
- 研究に協力する人は必ず「研究対象者」または「参加者」と表記し、旧来の呼称は使わない

見出しは含めず、本文のみを出力してください。
"""
        result = await self._call_llm(prompt)
        return result if result.strip() else self._fallback_overview(context)

    async def _generate_experiment_objective(self, context: Dict[str, Any], is_questionnaire: bool = False) -> str:
        """実験の目的 / アンケートの目的を生成（3-1）"""
        if is_questionnaire:
            prompt = f"""
{IMPLEMENTATION_PLAN_RULES}

【タスク】
公式参考様式の「3-1.アンケートの目的」セクションを執筆してください。

研究タイトル: {context['research_title']}
研究概要: {context['brief_description']}
研究方法: {context['methodology']}

【記述すべき内容】
1. このアンケート調査で何を把握・測定するか
2. どのような項目（態度、経験、評価など）を尋ねるか
3. 得られた回答を何の分析に用いるか

【執筆スタイル】
- 「本アンケートは〜を把握することを目的とする」のような断定形
- 1-2段落（100-200文字）で記述
- 段落間の空行は入れない

見出しは含めず、本文のみを出力してください。
"""
        else:
            prompt = f"""
{IMPLEMENTATION_PLAN_RULES}

【タスク】
公式参考様式の「3-1.実験の目的」セクションを執筆してください。
これは研究全体の目的とは異なり、具体的な実験手続きの目的を説明するものです。

研究タイトル: {context['research_title']}
研究方法: {context['methodology']}
使用機器: {', '.join(context['devices']) if context['devices'] else '特になし'}

【記述すべき内容】
1. この実験で何を測定・観察するか
2. どのような条件を設定し比較するか
3. 参加者にどのような回答・行動を求めるか

【執筆スタイル】
- 「本実験は〜を測定することで〜を明らかにすることを目的とする」のような形式
- 1段落（100-150文字）で簡潔に記述

見出しは含めず、本文のみを出力してください。
"""
        result = await self._call_llm(prompt)
        return result if result.strip() else self._fallback_experiment_objective(context)

    async def _generate_participants(self, context: Dict[str, Any], is_questionnaire: bool = False) -> str:
        """実験参加者 / 研究対象者を生成（3-2）。同意書裏面相当の厚みを持たせる。"""
        section_label = "研究対象者" if is_questionnaire else "実験参加者"
        section_number = "3-2.研究対象者" if is_questionnaire else "3-2.実験参加者"
        prompt = f"""
{IMPLEMENTATION_PLAN_RULES}

【タスク】
公式参考様式の「{section_number}」セクションを執筆してください。

研究タイトル: {context['research_title']}
研究対象者の概要: {context['target_participants']}
予定参加者数: {context['participant_count']}名
想定されるリスク: {', '.join(context['risks']) if context['risks'] else '特になし'}

【記述すべき内容（順番通りに、複数文で具体的に）】
1. {section_label}の条件（年齢、所属、視力・聴力など研究に関係する要件）を「（1）」「（2）」のように列挙し、予定人数を明記する
2. この研究にこれらの研究対象者が必要である理由（必要性）を簡潔に述べる
3. 募集方法（WEB公募、学内掲示、研究室広報など）と、別紙の募集文を用いることがあれば言及する
4. 謝金は大学の謝金規程に基づき支払うこと
5. 参加は自由意思によること、参加しない場合や途中で取りやめた場合にも不利益がないこと

【執筆スタイル】
- 「{section_label}は、（1）〜、（2）〜の条件を満たす者〇〇名とする。」のような断定形で開始する
- 公式参考様式の記載例に倣い、簡潔だが必要事項を網羅した複数文で記述する
- 段落間の空行は入れない
- 研究に協力する人は必ず「研究対象者」または「参加者」と表記し、旧来の呼称は使わない

見出しは含めず、本文のみを出力してください。
"""
        result = await self._call_llm(prompt)
        return result if result.strip() else self._fallback_participants(context)

    async def _generate_reward(self, context: Dict[str, Any]) -> str:
        """謝金についてを生成"""
        duration = context.get('duration', 60)
        reward = context.get('reward_amount', 1230)
        
        prompt = f"""
{IMPLEMENTATION_PLAN_RULES}

【タスク】
「謝金について」セクションを執筆してください。

所要時間: 約{duration}分
謝金額: {reward}円

【記述すべき内容】
1. 謝礼の形式と金額
2. 算出根拠（最低賃金または大学規定に基づく計算根拠）

【参考フォーマット】
本研究の参加者への謝礼として、Amazonギフトカード（Eメールタイプ）{reward}円分を配布する。
【算出根拠】本学の規定に基づき、実験所要時間（{duration}分）相当額を算出した。

1段落で簡潔に記述してください。
見出しは含めず、本文のみを出力してください。
"""
        return await self._call_llm(prompt)
    
    async def _generate_equipment(self, context: Dict[str, Any]) -> str:
        """実験装置・実験タスクを生成（3-3）"""
        devices = context.get('devices', [])
        devices_str = ', '.join(devices) if devices else '特になし'

        prompt = f"""
{IMPLEMENTATION_PLAN_RULES}

【タスク】
公式参考様式の「3-3.実験装置・実験タスク」セクションを執筆してください。

研究方法: {context['methodology']}
使用機器: {devices_str}

【記述すべき内容】
1. 実験で使用する装置・機器を「第一に」「第二に」と列挙して説明する
   - 各装置の仕様や役割を具体的に記述
   - 安全性に関わる配置や設定があれば記述
2. 参加者が取り組む実験タスクの内容（何を提示し、何に回答・反応してもらうか）を具体的に記述する

【執筆スタイル】
- 「本実験で用いる装置は以下から構成される。」で開始
- 各装置を「第一に、〜である。」「第二に、〜である。」の形式で説明
- 続けて実験タスクを「実験タスクとして、参加者は〜する。」のように記述
- 段落間の空行は入れない

見出しは含めず、本文のみを出力してください。
"""
        result = await self._call_llm(prompt)
        return result if result.strip() else self._fallback_equipment(context)

    async def _generate_procedures(self, context: Dict[str, Any], is_questionnaire: bool = False) -> str:
        """実験手順（3-4）/ 実施内容（アンケートの場合 3-3）を生成"""
        duration = context.get('duration', 60)
        if is_questionnaire:
            prompt = f"""
{IMPLEMENTATION_PLAN_RULES}

【タスク】
公式参考様式の「3-3.実施内容」セクションを執筆してください。

研究方法: {context['methodology']}
所要時間: 約{duration}分

【記述すべき内容】
1. アンケートの実施方法（WEBフォーム、紙の質問紙など）と実施の流れ
2. 研究内容について書面で説明し、同意を得た上で実施すること
3. 回答に要するおおよその時間
4. 参加は任意であり、回答を望まない設問には答えなくてよいこと、いつでも中止できること

【執筆スタイル】
- 「本アンケートに関する説明を書面で行った上で、〜を実施する。」のような断定形で開始
- 段落間の空行は入れない
- 研究に協力する人は必ず「研究対象者」または「参加者」と表記し、旧来の呼称は使わない

見出しは含めず、本文のみを出力してください。
"""
        else:
            prompt = f"""
{IMPLEMENTATION_PLAN_RULES}

【タスク】
公式参考様式の「3-4.実験手順」セクションを執筆してください。

研究方法: {context['methodology']}
所要時間: 約{duration}分
使用機器: {', '.join(context['devices']) if context['devices'] else '特になし'}

【記述すべき手順（公式記載例に倣い「本実験に関する説明を書面で行った上で、…」から始める）】
1. 研究の目的・内容・倫理的配慮の説明と同意取得（約10分）
   - 何を説明するか、参加者の権利をどう伝えるか
2. 実験準備および姿勢の調整（約5分）
   - 機器の設定、参加者の準備
3. 実験試行（約X分 × 条件数）
   - 各試行で何を行うか、休憩の有無
4. 機器の取り外し・終了処理（約5分）
   - 終了確認、体調確認

【執筆スタイル】
- 「本実験に関する説明を書面で行った上で、」で始める
- 各手順のサブ見出しは太字で「手順名（約X分）」の形式
- 実験を実施する担当者は「実験実施者」と表記し、実験実施者と参加者の行動を具体的に記述する
- いつでも中断可能であること、不快時の対応を明記
- 段落間の空行は入れない
- 研究に協力する人は必ず「研究対象者」または「参加者」と表記し、旧来の呼称は使わない

見出しは含めず、本文のみを出力してください。
"""
        result = await self._call_llm(prompt)
        return result if result.strip() else self._fallback_procedures(context)

    async def _generate_risks(self, context: Dict[str, Any]) -> str:
        """想定される負荷を生成"""
        risks = context.get('risks', [])
        countermeasures = context.get('risk_countermeasures', [])
        risks_str = '\n'.join(f"- {r}" for r in risks) if risks else "特になし"
        countermeasures_str = '\n'.join(f"- {c}" for c in countermeasures) if countermeasures else ""
        
        prompt = f"""
{IMPLEMENTATION_PLAN_RULES}

【タスク】
「想定される精神的・物理的負荷」セクションを執筆してください。

研究方法: {context['methodology']}
使用機器: {', '.join(context['devices']) if context['devices'] else '特になし'}
想定されるリスク:
{risks_str}
対策:
{countermeasures_str}

【記述すべき内容】
1. 精神的負荷
   - どのような心理的負担が生じうるか
   - 対処法（事前説明、中断可能性など）
2. 物理的負荷
   - 「第一に」「第二に」と列挙してリスクを説明
   - 各リスクへの対処法
3. 実験装置の安全性およびリスク管理について（見出し付き）
   - 常時監視体制
   - 機器配置による安全確保

【執筆スタイル】
- 「本研究では、以下の精神的負荷が生じる可能性がある。」で開始
- 具体的な対策を必ず記述
- 段落間の空行は入れない

見出しは含めず、本文のみを出力してください。
"""
        return await self._call_llm(prompt)
    
    # ------------------------------------------------------------------
    # 決定的フォールバック（LLMが空応答のとき、context から本文を組み立てて
    # セクションを空欄にしない。中断せず必ずドラフトを出す方針に沿う）。
    # ------------------------------------------------------------------
    @staticmethod
    def _ordinal_jp(n: int) -> str:
        kanji = "一二三四五六七八九十"
        return kanji[n - 1] if 1 <= n <= len(kanji) else str(n)

    def _fallback_overview(self, context: Dict[str, Any]) -> str:
        purpose = self._as_sentence(context.get("brief_description", ""))
        method = self._as_sentence(context.get("methodology", ""))
        target_raw = str(context.get("target_participants", "")).strip().rstrip("。．.").strip()
        parts = [p for p in (purpose, method) if p]
        if target_raw:
            parts.append(f"研究対象者は{target_raw}とする。")
        return "".join(parts) or "本研究の概要は別紙のとおりである。"

    @staticmethod
    def _as_sentence(text: str) -> str:
        """文字列を1文として整える（前後空白を除き、末尾に句点を付ける）。空なら空文字。"""
        t = str(text).strip()
        if not t:
            return ""
        return t if t.endswith(("。", "．", ".")) else t + "。"

    def _fallback_experiment_objective(self, context: Dict[str, Any]) -> str:
        # 目的は purpose を優先（無ければ method）。全文を文中に埋め込まず独立した文として並べ、
        # 「本実験は、本実験は…」のような主語重複や「…。を…」の句読点崩れを避ける。
        purpose = self._as_sentence(context.get("brief_description", ""))
        method = self._as_sentence(context.get("methodology", ""))
        parts = []
        if purpose:
            parts.append(purpose)
        elif method:
            parts.append(method)
        parts.append("本実験では、設定した条件間で参加者の反応や回答を比較し、研究目的に関わる指標を測定する。")
        return "".join(parts)

    def _fallback_participants(self, context: Dict[str, Any]) -> str:
        target = self._as_sentence(context.get("target_participants", ""))
        count = context.get("participant_count", "")
        count_str = f"{count}名" if count not in (None, "", 0) else "所定の人数"
        parts = []
        if target:
            parts.append(target)
            parts.append(f"予定人数は{count_str}とする。")
        else:
            parts.append(f"実験参加者は本研究の参加条件を満たす成人とし、予定人数は{count_str}とする。")
        parts.append("本研究の目的を達成するため、これらの参加者の反応や回答を分析する必要がある。")
        parts.append(
            "募集は学内掲示およびメール・SNS等による公募で行い、研究室内で募集する場合は、"
            "参加・不参加が成績評価や指導上の関係に影響しないことを明示し、参加の自由意思を担保する。"
        )
        parts.append("謝金は大学の謝金規程に基づき支払う。")
        parts.append("参加は自由意思によるものであり、参加しない場合や途中で取りやめた場合にも不利益は生じない。")
        return "".join(parts)

    def _fallback_equipment(self, context: Dict[str, Any]) -> str:
        devices = [str(d).strip() for d in (context.get("devices") or []) if str(d).strip()]
        method = self._as_sentence(context.get("methodology", ""))
        if devices:
            sentences = "".join(
                f"第{self._ordinal_jp(i)}に、{device}を用いる。"
                for i, device in enumerate(devices, 1)
            )
            lead = f"本実験で用いる装置は以下から構成される。{sentences}"
        else:
            lead = (
                "本実験で用いる装置は以下から構成される。第一に、参加者が操作する個人用パソコン"
                "（またはノートパソコン）である。第二に、音声を提示するためのヘッドホンまたは"
                "イヤホンである。これらは一般的な機器であり、参加者に過度な負担を与えない。"
            )
        # タスクは method を独立した文として記述（文中に埋め込まない）
        if method:
            task = "実験タスクは次のとおりである。" + method
        else:
            task = "実験タスクとして、参加者は提示される刺激を視聴し、所定の課題に回答する。"
        return lead + task

    def _fallback_procedures(self, context: Dict[str, Any]) -> str:
        duration = context.get("duration", 60) or 60
        method = str(context.get("methodology", "")).strip()
        return (
            "本実験に関する説明を書面で行った上で、以下の手順で実施する。"
            "第一に、研究の目的・内容・倫理的配慮について説明し、参加の任意性といつでも中断・撤回できる"
            "ことを伝えた上で同意を取得する（約10分）。第二に、使用機器の準備と動作確認を行う（約5分）。"
            f"第三に、{method or '設定した条件に基づく課題を提示し、参加者は提示内容に回答する'}"
            "（中心となる実験試行）。各試行の合間には適宜休憩を挟み、参加者はいつでも中断できる。"
            "第四に、終了処理として体調を確認し、データの保存と謝礼の案内を行う（約5分）。"
            f"全体の所要時間は約{duration}分である。"
        )

    async def _call_llm(self, prompt: str) -> str:
        """LLM呼び出し（共通処理）。

        推論モデルが一過性に空本文を返すことがあるため、空のときは一度だけ再試行する。
        最終的に空ならば空文字を返し、呼び出し側が決定的フォールバックで埋める。
        """
        system_instruction = (
            "あなたは日本の大学で研究倫理審査申請書を作成する経験豊富な研究者です。"
            "実施計画書は「これから行う」計画ではなく「決定済み」の内容を記述するものです。"
            "すべて断定形で書き、曖昧な表現は避けてください。"
            "専門用語は使用せず、一般の方にも分かりやすい日本語で記述してください。"
            "である調で統一してください。"
            + TERMINOLOGY_RULE
        )
        for attempt in range(2):
            try:
                response = await self.llm.generate_content_async(
                    prompt=prompt,
                    system_instruction=system_instruction,
                )
                processed = self._post_process(response)
                if processed.strip():
                    return processed
                logger.warning(f"LLMが空応答を返しました（再試行 {attempt + 1}/2）")
            except Exception as e:
                logger.error(f"LLM呼び出しエラー（再試行 {attempt + 1}/2）: {e}")
        return ""
    
    def _post_process(self, text: str) -> str:
        """AIライクな表現を除去する後処理"""
        if not text:
            return ""
        
        # 不要な記号を削除
        replacements = {
            ":": "。",
            "!": "。",
            "~~": "",
            "**": "",
            "##": "",
            "###": "",
        }
        
        for old, new in replacements.items():
            text = text.replace(old, new)
        
        # 断定的表現を柔らかく→逆に断定を維持
        academic_replacements = {
            "素晴らしい": "",
            "ことができます": "ことが可能である",
            "を行います": "を行う",
            "と思われます": "である",
            "かもしれません": "である",
            "と考えられます": "である",
        }
        
        for old, new in academic_replacements.items():
            text = text.replace(old, new)
        
        # 連続する空行を削除
        while "\n\n\n" in text:
            text = text.replace("\n\n\n", "\n\n")

        # 用語統一（被験者/健常者→研究対象者・参加者、研究者→実験実施者）。
        # LLM 生成物にもプロンプト指示だけでは残ることがあるため決定的に置換する。
        from app.services.context_text_enricher import normalize_research_terminology
        text = normalize_research_terminology(text)

        return text.strip()
    
    def _build_implementation_plan_docx(
        self,
        sections: Dict[str, str],
        context: Dict[str, Any],
        outline: List[Dict[str, Any]],
        output_dir: Path
    ) -> Path:
        """実施計画書DOCXを公式参考様式 260127 の章立てで構築する。"""
        doc = Document()

        # スタイル設定
        style = doc.styles['Normal']
        style.font.name = 'Yu Gothic'
        style.font.size = Pt(11)

        # タイトル
        title = doc.add_paragraph()
        title_run = title.add_run('実施計画書')
        title_run.bold = True
        title_run.font.size = Pt(14)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()

        # outline に従って各セクションを描画
        for item in outline:
            number = item["number"]
            label = f"{number}.{item['title']}"
            if item["level"] == 1:
                self._add_heading(doc, label)
            else:
                self._add_subheading(doc, label)

            key = item.get("key")
            if key == "research_title":
                doc.add_paragraph(context.get('research_title', ''))
            elif key:
                self._add_text_no_blank(doc, sections.get(key, ''))

        # 保存
        output_path = output_dir / "実施計画書.docx"
        doc.save(output_path)

        return output_path

    def _add_heading(self, doc: Document, text: str):
        """大見出しを追加（太字、前に空行）"""
        # 見出し前に空行を追加
        doc.add_paragraph()
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(12)
        # 段落の後に少しスペース
        p.paragraph_format.space_after = Pt(6)

    def _add_subheading(self, doc: Document, text: str):
        """小見出しを追加（太字、わずかにインデント）"""
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(11)
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(3)
    
    def _add_text_no_blank(self, doc: Document, text: str):
        """テキストを追加（適切な段落間隔）"""
        if not text:
            return
        
        # 改行で分割
        paragraphs = text.split('\n')
        for para_text in paragraphs:
            para_text = para_text.strip()
            if para_text:
                p = doc.add_paragraph()
                
                # 「対策:」で始まる行はインデント
                if para_text.startswith("対策:") or para_text.startswith("　対策:") or para_text.startswith("  対策:"):
                    p.paragraph_format.left_indent = Cm(1)
                    p.paragraph_format.first_line_indent = Cm(-0.5)
                
                run = p.add_run(para_text)
                run.font.size = Pt(10.5)
                
                # 行間を少し広げて読みやすく
                p.paragraph_format.line_spacing = 1.3
                p.paragraph_format.space_after = Pt(3)


async def generate_implementation_plan_with_llm(
    form_data: Dict[str, Any],
    output_dir: Path,
    llm_client: LLMClient,
    lab_defaults: Dict[str, Any]
) -> Path:
    """
    実施計画書をLLMで生成（エントリーポイント）
    """
    generator = LLMDocumentGenerator(llm_client, lab_defaults)
    return await generator.generate_implementation_plan(form_data, output_dir)
